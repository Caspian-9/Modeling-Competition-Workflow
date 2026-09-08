from __future__ import annotations

import csv
import os
import re
import shutil
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from .models import ValidationError, hash_file, utc_now
from .storage import CaseRepository


PROBLEM_SUFFIXES = frozenset({".pdf", ".md", ".txt"})
DATA_SUFFIXES = frozenset({".xlsx", ".xls", ".csv", ".tsv"})


class RawIntakeService:
    """Inventory raw files and extract problem text without a central artifact registry."""

    def __init__(self, repository: CaseRepository) -> None:
        self.repository = repository

    def ingest(
        self,
        case_id: str,
        source_paths: list[str | Path],
        *,
        replace: bool = False,
    ) -> dict[str, Any]:
        """Copy chat attachments into raw/ before running normal intake."""
        config = self.repository.load_case(case_id)
        if not source_paths:
            raise ValidationError("chat intake requires at least one attached file")
        case_dir = self.repository.case_dir(config.case_id)
        raw_dir = case_dir / "raw"
        copied: list[dict[str, Any]] = []
        seen_names: set[str] = set()
        for index, value in enumerate(source_paths):
            source = Path(value).expanduser().resolve()
            if not source.is_file():
                raise ValidationError(f"attachment[{index}] is not a readable file: {value}")
            if source.name.lower() in seen_names:
                raise ValidationError(f"duplicate attachment filename: {source.name}")
            seen_names.add(source.name.lower())
            suffix = source.suffix.lower()
            if suffix not in PROBLEM_SUFFIXES | DATA_SUFFIXES:
                raise ValidationError(
                    f"unsupported attachment type: {source.name}; "
                    "attach a PDF/text problem or XLSX/XLS/CSV/TSV data file"
                )
            destination = raw_dir / source.name
            copied.append(
                self._copy_attachment(source, destination, replace=replace)
            )
        return {"copied_files": copied, "inventory": self.sync(config.case_id)}

    def sync(self, case_id: str) -> dict[str, Any]:
        config = self.repository.load_case(case_id)
        case_dir = self.repository.case_dir(config.case_id)
        raw_dir = case_dir / "raw"
        files = tuple(
            path
            for path in sorted(raw_dir.iterdir(), key=lambda item: item.name.lower())
            if path.is_file() and not path.name.startswith("~$")
        )
        problem_files = tuple(
            path for path in files if path.suffix.lower() in PROBLEM_SUFFIXES
        )
        data_files = tuple(path for path in files if path.suffix.lower() in DATA_SUFFIXES)
        unsupported = tuple(
            path.name
            for path in files
            if path.suffix.lower() not in PROBLEM_SUFFIXES | DATA_SUFFIXES
        )
        errors: list[str] = []
        warnings: list[str] = []
        if len(problem_files) != 1:
            errors.append(
                f"raw/ 必须恰好包含一份题目 PDF、Markdown 或文本，当前为 {len(problem_files)} 份"
            )
        if not data_files:
            errors.append("raw/ 至少需要一个 .xlsx、.xls、.csv 或 .tsv 数据文件")
        if unsupported:
            warnings.append("忽略不支持的文件：" + "、".join(unsupported))

        problem: dict[str, Any] | None = None
        extraction: dict[str, Any] = {
            "status": "NOT_READY",
            "path": None,
            "message": None,
        }
        if len(problem_files) == 1:
            problem = self._file_record(problem_files[0], case_dir)
            extraction = self._extract_problem(problem_files[0], case_dir)
            if extraction["status"] != "READY":
                errors.append(str(extraction["message"]))

        data_records: list[dict[str, Any]] = []
        derived_csv_files: list[dict[str, Any]] = []
        for path in data_files:
            record = self._file_record(path, case_dir)
            record["metadata"] = self._data_metadata(path)
            inspection_error = record["metadata"].get("inspection_error")
            if inspection_error:
                errors.append(f"数据文件 {path.name} 无法检查：{inspection_error}")
            elif path.suffix.lower() == ".xlsx":
                try:
                    conversions = self._convert_xlsx_to_csv(path, case_dir)
                    record["metadata"]["csv_exports"] = conversions
                    derived_csv_files.extend(conversions)
                except ValidationError as exc:
                    errors.append(f"数据文件 {path.name} 无法转换为 CSV：{exc}")
            data_records.append(record)

        inventory_path = case_dir / "workspace" / "raw_inventory.json"
        previous = (
            self.repository.read_json(inventory_path) if inventory_path.is_file() else None
        )
        changed = self._changed_paths(previous, problem, data_records)
        if changed:
            warnings.append(
                "以下 raw 文件相对上次登记已变化，请人工复核相关分析："
                + "、".join(changed)
            )

        inventory = {
            "schema_version": "fast-1.0",
            "case_id": config.case_id,
            "scanned_at": utc_now(),
            "ready": not errors,
            "problem": problem,
            "problem_extraction": extraction,
            "data_files": data_records,
            "derived_csv_files": derived_csv_files,
            "errors": errors,
            "warnings": warnings,
        }
        self.repository.write_json(inventory_path, inventory)
        return inventory

    @staticmethod
    def _copy_attachment(
        source: Path, destination: Path, *, replace: bool
    ) -> dict[str, Any]:
        source_hash = hash_file(source)
        if destination.exists():
            if destination.resolve() == source.resolve():
                action = "already_in_raw"
            elif hash_file(destination) == source_hash:
                action = "reused_identical"
            elif not replace:
                raise ValidationError(
                    f"raw/{destination.name} already exists with different content; "
                    "use --replace after confirming the new attachment is intended"
                )
            else:
                RawIntakeService._atomic_copy(source, destination)
                action = "replaced"
        else:
            RawIntakeService._atomic_copy(source, destination)
            action = "copied"
        return {
            "source_path": str(source),
            "raw_path": destination.name,
            "sha256": source_hash,
            "action": action,
        }

    @staticmethod
    def _atomic_copy(source: Path, destination: Path) -> None:
        destination.parent.mkdir(parents=True, exist_ok=True)
        fd, temporary = tempfile.mkstemp(
            prefix=f".{destination.name}.", suffix=".tmp", dir=destination.parent
        )
        os.close(fd)
        try:
            shutil.copyfile(source, temporary)
            shutil.copystat(source, temporary)
            os.replace(temporary, destination)
        finally:
            temporary_path = Path(temporary)
            if temporary_path.exists():
                temporary_path.unlink()

    @staticmethod
    def _file_record(path: Path, case_dir: Path) -> dict[str, Any]:
        return {
            "path": path.relative_to(case_dir).as_posix(),
            "sha256": hash_file(path),
            "size_bytes": path.stat().st_size,
            "format": path.suffix.lower().lstrip("."),
        }

    def _extract_problem(self, source: Path, case_dir: Path) -> dict[str, Any]:
        suffix = source.suffix.lower()
        if suffix in {".md", ".txt"}:
            try:
                text = source.read_text(encoding="utf-8-sig").strip()
            except UnicodeDecodeError:
                text = source.read_text(encoding="gb18030").strip()
            page_count = None
        else:
            try:
                from pypdf import PdfReader
            except ImportError:
                return {
                    "status": "BLOCKED",
                    "path": None,
                    "message": "无法提取题面：缺少可选依赖 pypdf",
                }
            try:
                reader = PdfReader(str(source))
                pages = [(page.extract_text() or "").strip() for page in reader.pages]
            except Exception as exc:
                return {
                    "status": "BLOCKED",
                    "path": None,
                    "message": f"题目 PDF 文本提取失败：{type(exc).__name__}: {exc}",
                }
            text = "\n\n".join(page for page in pages if page).strip()
            page_count = len(pages)
        if len(text) < 20:
            return {
                "status": "BLOCKED",
                "path": None,
                "message": "题面未提取到足够文本；扫描版 PDF 需要先 OCR",
            }
        target = case_dir / "workspace" / "problem.md"
        target.write_text(
            f"# 题目\n\n来源：`raw/{source.name}`\n\n{text}\n",
            encoding="utf-8",
        )
        result: dict[str, Any] = {
            "status": "READY",
            "path": target.relative_to(case_dir).as_posix(),
            "message": None,
        }
        if page_count is not None:
            result["page_count"] = page_count
        return result

    def _convert_xlsx_to_csv(self, source: Path, case_dir: Path) -> list[dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValidationError(
                "requires openpyxl; install it in the desktop workspace before retrying"
            ) from exc
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
        except Exception as exc:
            raise ValidationError(f"{type(exc).__name__}: {exc}") from exc

        source_hash = hash_file(source)
        target_root = (
            case_dir
            / "data"
            / "interim"
            / "raw_csv"
            / f"{self._safe_filename(source.stem)}-{source_hash[:12]}"
        )
        target_root.mkdir(parents=True, exist_ok=True)
        records: list[dict[str, Any]] = []
        try:
            for index, worksheet in enumerate(workbook.worksheets, start=1):
                target = target_root / (
                    f"{index:02d}-{self._safe_filename(worksheet.title)}.csv"
                )
                row_count = 0
                with target.open("w", encoding="utf-8", newline="") as handle:
                    writer = csv.writer(handle, lineterminator="\n")
                    for row in worksheet.iter_rows(values_only=True):
                        writer.writerow([self._csv_value(value) for value in row])
                        row_count += 1
                records.append(
                    {
                        "source_path": source.relative_to(case_dir).as_posix(),
                        "source_sha256": source_hash,
                        "sheet_name": worksheet.title,
                        "path": target.relative_to(case_dir).as_posix(),
                        "sha256": hash_file(target),
                        "row_count": row_count,
                        "column_count": worksheet.max_column,
                        "encoding": "utf-8",
                        "formula_mode": "cached_values",
                    }
                )
        finally:
            workbook.close()
        return records

    @staticmethod
    def _safe_filename(value: str) -> str:
        cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip(".-")
        return cleaned[:64] or "sheet"

    @staticmethod
    def _csv_value(value: Any) -> str | int | float:
        if value is None:
            return ""
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return value

    @staticmethod
    def _changed_paths(
        previous: Any,
        problem: dict[str, Any] | None,
        data_files: list[dict[str, Any]],
    ) -> list[str]:
        if not isinstance(previous, dict):
            return []
        old_records = [
            item
            for item in [previous.get("problem"), *previous.get("data_files", [])]
            if isinstance(item, dict) and isinstance(item.get("path"), str)
        ]
        new_records = [
            item for item in [problem, *data_files] if isinstance(item, dict)
        ]
        old = {item["path"]: item.get("sha256") for item in old_records}
        new = {item["path"]: item.get("sha256") for item in new_records}
        return sorted(path for path in old.keys() | new.keys() if old.get(path) != new.get(path))

    @staticmethod
    def _data_metadata(path: Path) -> dict[str, Any]:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            try:
                with ZipFile(path) as archive:
                    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
                    namespace = {
                        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                    }
                    sheets = [
                        item.attrib.get("name", "")
                        for item in workbook.findall("m:sheets/m:sheet", namespace)
                    ]
                return {"sheet_names": sheets, "sheet_count": len(sheets)}
            except (BadZipFile, KeyError, ElementTree.ParseError) as exc:
                return {"inspection_error": f"{type(exc).__name__}: {exc}"}
        if suffix in {".csv", ".tsv"}:
            delimiter = "\t" if suffix == ".tsv" else ","
            try:
                with path.open("r", encoding="utf-8-sig", newline="") as handle:
                    header = next(csv.reader(handle, delimiter=delimiter), [])
                return {"columns": header, "column_count": len(header)}
            except (OSError, UnicodeError, csv.Error) as exc:
                return {"inspection_error": f"{type(exc).__name__}: {exc}"}
        if suffix == ".xls":
            return {"inspection": "legacy XLS; schema inspection deferred to Engineer"}
        raise ValidationError(f"unsupported data format: {suffix}")
